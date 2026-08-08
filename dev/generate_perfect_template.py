import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()

# Font Styles
title_font = Font(name='Calibri', bold=True, size=16, color='1E3A8A')
section_font = Font(name='Calibri', bold=True, size=12, color='1E3A8A')
hdr_font = Font(name='Calibri', bold=True, size=11, color='FFFFFF')
body_font = Font(name='Calibri', size=10)
bold_body = Font(name='Calibri', bold=True, size=10)
italic_note = Font(name='Calibri', size=10, italic=True, color='4B5563')
code_font = Font(name='Consolas', size=10, bold=True, color='B91C1C')

# Fills
hdr_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid') # Navy
sample_fill = PatternFill(start_color='FEF3C7', end_color='FEF3C7', fill_type='solid') # Light Yellow
sec_fill = PatternFill(start_color='DBEAFE', end_color='DBEAFE', fill_type='solid') # Soft Blue

# Borders
thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1'),
)

# Alignments
hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
valign_top = Alignment(vertical='top', wrap_text=True)

# ============================================================
# SHEET 1: PETUNJUK PENGISIAN
# ============================================================
ws_inst = wb.active
ws_inst.title = 'Petunjuk Pengisian'
ws_inst.sheet_properties.tabColor = '1E3A8A'

ws_inst['A1'].value = 'PANDUAN & PETUNJUK PENGISIAN TEMPLATE IMPORT SOAL'
ws_inst['A1'].font = title_font

ws_inst['A2'].value = 'Platform YakinLulus.id — Format Standar Import Excel (.xlsx)'
ws_inst['A2'].font = italic_note

instructions_sections = [
    ("1. ATURAN UMUM", [
        ("Nama Sheet Data", "Wajib bernama 'Soal' atau 'Soal (Isi)'. Jangan mengubah nama sheet ini."),
        ("Header Tabel", "Baris ke-1 adalah Header. Jangan menghapus, menggeser, atau mengubah teks Header."),
        ("Baris Data", "Data soal diisikan mulai dari Baris ke-2 ke bawah. Setiap baris mewakili 1 Soal."),
        ("Mapel & Konten", "Kolom 'Mapel' dan 'Blok 1 - Isi' WAJIB diisi. Baris dengan Mapel kosong akan diabaikan."),
    ]),
    ("2. PENJELASAN TIPE SOAL & FORMAT KUNCI JAWABAN", [
        ("SINGLE_CHOICE", "Pilihan Ganda biasa (1 jawaban benar). Kunci Jawaban diisi 1 huruf saja: A, B, C, D, atau E. Opsi A..E diisi pilihan jawaban."),
        ("MULTIPLE_CHOICE", "Pilihan Ganda Kompleks (>1 jawaban benar). Kunci Jawaban diisi huruf dipisahkan koma atau digabung: A,C atau A,B,D. Opsi A..E diisi pilihan jawaban."),
        ("TRUE_FALSE", "Matriks Benar - Salah. Opsi A, B, C, D diisi teks PERNYATAAN. Kunci Jawaban diisi format: A:B, B:S, C:B, D:S (atau B,S,B,S)."),
    ]),
    ("3. DESKRIPSI RINCI KOLOM", [
        ("Mapel", "Nama/Kode Mata Pelajaran (contoh: 'Matematika Wajib Kelas 10' atau 'MTK-WAJIB-10'). Harus ada di database."),
        ("Kelas", "Tingkat kelas (contoh: 10, 11, 12, atau UTBK)."),
        ("Bab (ID)", "ID Bab (UUID) jika ada, atau dapat dikosongkan jika belum dikelompokkan ke Bab."),
        ("Tipe Soal", "Pilih: SINGLE_CHOICE, MULTIPLE_CHOICE, atau TRUE_FALSE."),
        ("Kesulitan", "Pilih: EASY, MEDIUM, atau HARD (default: MEDIUM)."),
        ("Blok 1..4", "Blok 1 - Tipe: PARAGRAPH (teks) atau IMAGE (URL/ID Gambar). Blok 1 - Isi: Teks Utama Soal."),
        ("Opsi A s/d H", "Teks Pilihan Jawaban (untuk PG/PGK) atau Teks Pernyataan (untuk True/False). Kolom opsi yang tidak terpakai dikosongkan."),
        ("Kunci Jawaban", "Huruf jawaban benar atau matriks Benar/Salah (Lihat Aturan Bagian 2)."),
        ("Skor / Skor Negatif", "Skor jika benar (default: 1.0 atau 5.0) dan skor negatif jika salah (default: 0.0)."),
        ("Pembahasan", "Teks atau penjelasan lengkap cara menyelesaikan soal."),
        ("Bloom Level & Bahasa", "Bloom Level: C1, C2, C3, C4, C5, C6. Bahasa: 'id' (Indonesia) atau 'en' (Inggris)."),
    ])
]

curr_row = 4
for sec_title, items in instructions_sections:
    ws_inst.cell(row=curr_row, column=1, value=sec_title).font = section_font
    ws_inst.cell(row=curr_row, column=1).fill = sec_fill
    ws_inst.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
    curr_row += 1
    
    for item_title, item_desc in items:
        c1 = ws_inst.cell(row=curr_row, column=1, value=item_title)
        c1.font = bold_body
        c1.alignment = valign_top
        
        c2 = ws_inst.cell(row=curr_row, column=2, value=item_desc)
        c2.font = body_font
        c2.alignment = valign_top
        curr_row += 1
    curr_row += 1

ws_inst.column_dimensions['A'].width = 24
ws_inst.column_dimensions['B'].width = 90

# ============================================================
# SHEET 2: SOAL (DATA TABLE)
# ============================================================
ws_data = wb.create_sheet('Soal')
ws_data.sheet_properties.tabColor = '10B981' # Green

headers = [
    "No", "Kode", "Mapel", "Kelas", "Bab (ID)", "Tipe Soal", "Kesulitan",
    "Blok 1 - Tipe", "Blok 1 - Isi",
    "Blok 2 - Tipe", "Blok 2 - Isi",
    "Blok 3 - Tipe", "Blok 3 - Isi",
    "Blok 4 - Tipe", "Blok 4 - Isi",
    "Opsi A", "Opsi B", "Opsi C", "Opsi D", "Opsi E", "Opsi F", "Opsi G", "Opsi H",
    "Kunci Jawaban", "Skor", "Skor Negatif", "Pembahasan", "Bloom Level", "Bahasa"
]

for col_idx, h in enumerate(headers, 1):
    cell = ws_data.cell(row=1, column=col_idx, value=h)
    cell.font = hdr_font
    cell.fill = hdr_fill
    cell.alignment = hdr_align
    cell.border = thin_border

# Sample Rows
sample_rows = [
    # Row 1: SINGLE_CHOICE
    [
        1, "SOAL-MTK-001", "Matematika", "10", "", "SINGLE_CHOICE", "MEDIUM",
        "PARAGRAPH", "Sebuah toko menjual 2 jenis buah: apel Rp 5.000/ons dan jeruk Rp 3.000/ons. Jika beli total 10 ons seharga Rp 42.000, maka banyak apel yang dibeli adalah ...",
        "", "", "", "", "", "",
        "6 ons", "4 ons", "5 ons", "3 ons", "8 ons", "", "", "",
        "A", 5, 0, "Misalkan x = apel, y = jeruk. x + y = 10 dan 5000x + 3000y = 42000 => x = 6.", "C3", "id"
    ],
    # Row 2: MULTIPLE_CHOICE
    [
        2, "SOAL-FIS-002", "Fisika", "11", "", "MULTIPLE_CHOICE", "HARD",
        "PARAGRAPH", "Manakah dari besaran-besaran berikut yang merupakan besaran turunan? (Pilih semua yang benar)",
        "", "", "", "", "", "",
        "Kecepatan", "Massa", "Gaya", "Panjang", "Energi", "", "", "",
        "A,C,E", 5, 0, "Besaran turunan meliputi Kecepatan, Gaya, dan Energi. Massa dan Panjang adalah besaran pokok.", "C4", "id"
    ],
    # Row 3: TRUE_FALSE
    [
        3, "SOAL-BIO-003", "Biologi", "10", "", "TRUE_FALSE", "MEDIUM",
        "PARAGRAPH", "Tentukan apakah setiap pernyataan mengenai sel berikut bernilai BENAR atau SALAH:",
        "", "", "", "", "", "",
        "Mitosondria berfungsi sebagai tempat respirasi seluler.",
        "Dinding sel ditemukan pada sel hewan.",
        "Ribosom berperan dalam sintesis protein.",
        "Membran sel bersifat impermiabel terhadap semua zat.",
        "", "", "", "",
        "A:B, B:S, C:B, D:S", 5, 0, "Dinding sel hanya ada pada tumbuhan. Membran sel bersifat semi-permiabel.", "C3", "id"
    ]
]

for row_idx, rdata in enumerate(sample_rows, 2):
    for col_idx, val in enumerate(rdata, 1):
        cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
        cell.font = body_font
        cell.fill = sample_fill
        cell.border = thin_border
        cell.alignment = valign_top

# Blank rows with thin border up to row 100
for row_idx in range(5, 101):
    for col_idx in range(1, len(headers) + 1):
        cell = ws_data.cell(row=row_idx, column=col_idx)
        cell.font = body_font
        cell.border = thin_border
        cell.alignment = valign_top
    ws_data.cell(row=row_idx, column=1, value=f'=IF(C{row_idx}="","",ROW()-1)')

# Data validations
dv_tipe = DataValidation(type='list', formula1='"SINGLE_CHOICE,MULTIPLE_CHOICE,TRUE_FALSE"', allow_blank=True)
dv_tipe.error = 'Pilih tipe soal yang valid'
ws_data.add_data_validation(dv_tipe)
dv_tipe.add('F2:F100')

dv_diff = DataValidation(type='list', formula1='"EASY,MEDIUM,HARD"', allow_blank=True)
dv_diff.error = 'Pilih tingkat kesulitan'
ws_data.add_data_validation(dv_diff)
dv_diff.add('G2:G100')

dv_bloom = DataValidation(type='list', formula1='"C1,C2,C3,C4,C5,C6"', allow_blank=True)
ws_data.add_data_validation(dv_bloom)
dv_bloom.add('AB2:AB100')

dv_lang = DataValidation(type='list', formula1='"id,en"', allow_blank=True)
ws_data.add_data_validation(dv_lang)
dv_lang.add('AC2:AC100')

# Column Widths
col_widths = {
    "A": 6, "B": 15, "C": 22, "D": 10, "E": 24, "F": 18, "G": 12,
    "H": 14, "I": 45, "J": 14, "K": 30, "L": 14, "M": 30, "N": 14, "O": 30,
    "P": 30, "Q": 30, "R": 30, "S": 30, "T": 30, "U": 20, "V": 20, "W": 20,
    "X": 22, "Y": 10, "Z": 12, "AA": 40, "AB": 12, "AC": 10
}
for col_letter, w in col_widths.items():
    ws_data.column_dimensions[col_letter].width = w

ws_data.freeze_panes = 'A2'

# Save file
out_path = 'd:/Project/EdTech/Yakinlulus.id/template_import_soal.xlsx'
wb.save(out_path)
print(f"SUCCESS: Generated {out_path}")
